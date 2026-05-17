import io
from datetime import datetime, timedelta
from django.utils import timezone
from django.db.models import Count, Q
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, PageBreak
from reportlab.lib.enums import TA_CENTER, TA_LEFT
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.logs.models import Log
from apps.alerts.models import Alert

class ReportGenerator:
    
    def generate_pdf_report(self, start_date, end_date, report_type='summary'):
        """Generate PDF report"""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=72, leftMargin=72, topMargin=72, bottomMargin=18)
        
        elements = []
        styles = getSampleStyleSheet()
        
        # Title
        title_style = ParagraphStyle('CustomTitle', parent=styles['Heading1'], fontSize=24, textColor=colors.HexColor('#1a1a2e'), alignment=TA_CENTER, spaceAfter=30)
        elements.append(Paragraph('SecureWatch AI - Security Report', title_style))
        elements.append(Spacer(1, 12))
        
        # Report Info
        info_style = ParagraphStyle('Info', parent=styles['Normal'], fontSize=10, textColor=colors.grey)
        elements.append(Paragraph(f'Report Period: {start_date.strftime("%Y-%m-%d")} to {end_date.strftime("%Y-%m-%d")}', info_style))
        elements.append(Paragraph(f'Generated: {timezone.now().strftime("%Y-%m-%d %H:%M:%S")}', info_style))
        elements.append(Spacer(1, 20))
        
        # Statistics
        logs = Log.objects.filter(timestamp__range=[start_date, end_date])
        alerts = Alert.objects.filter(created_at__range=[start_date, end_date])
        
        stats_data = [
            ['Metric', 'Count'],
            ['Total Logs', logs.count()],
            ['Critical Logs', logs.filter(level='CRITICAL').count()],
            ['Error Logs', logs.filter(level='ERROR').count()],
            ['Threats Detected', logs.filter(ai_threat_detected=True).count()],
            ['Total Alerts', alerts.count()],
            ['Critical Alerts', alerts.filter(severity='critical').count()],
            ['Resolved Alerts', alerts.filter(status='resolved').count()],
        ]
        
        stats_table = Table(stats_data, colWidths=[3*inch, 2*inch])
        stats_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1a1a2e')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 12),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        elements.append(stats_table)
        elements.append(Spacer(1, 20))
        
        # Top Threats
        elements.append(Paragraph('Top Security Threats', styles['Heading2']))
        elements.append(Spacer(1, 12))
        
        threat_logs = logs.filter(ai_threat_detected=True).values('ai_threat_type').annotate(count=Count('id')).order_by('-count')[:5]
        
        if threat_logs:
            threat_data = [['Threat Type', 'Count']]
            for threat in threat_logs:
                threat_data.append([threat['ai_threat_type'] or 'Unknown', threat['count']])
            
            threat_table = Table(threat_data, colWidths=[3*inch, 2*inch])
            threat_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e74c3c')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            elements.append(threat_table)
        else:
            elements.append(Paragraph('No threats detected in this period.', styles['Normal']))
        
        elements.append(Spacer(1, 20))
        
        # Recent Critical Alerts
        elements.append(Paragraph('Recent Critical Alerts', styles['Heading2']))
        elements.append(Spacer(1, 12))
        
        critical_alerts = alerts.filter(severity='critical').order_by('-created_at')[:10]
        
        if critical_alerts:
            alert_data = [['Date', 'Title', 'Status']]
            for alert in critical_alerts:
                alert_data.append([
                    alert.created_at.strftime('%Y-%m-%d %H:%M'),
                    alert.title[:40] + '...' if len(alert.title) > 40 else alert.title,
                    alert.status.upper()
                ])
            
            alert_table = Table(alert_data, colWidths=[1.5*inch, 3*inch, 1*inch])
            alert_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#e74c3c')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, -1), 9),
                ('GRID', (0, 0), (-1, -1), 1, colors.black)
            ]))
            elements.append(alert_table)
        else:
            elements.append(Paragraph('No critical alerts in this period.', styles['Normal']))
        
        doc.build(elements)
        buffer.seek(0)
        return buffer
    
    def generate_excel_report(self, start_date, end_date):
        """Generate Excel report"""
        wb = openpyxl.Workbook()
        
        # Summary Sheet
        ws_summary = wb.active
        ws_summary.title = 'Summary'
        
        # Header
        ws_summary['A1'] = 'SecureWatch AI - Security Report'
        ws_summary['A1'].font = Font(size=16, bold=True, color='1a1a2e')
        ws_summary['A2'] = f'Period: {start_date.strftime("%Y-%m-%d")} to {end_date.strftime("%Y-%m-%d")}'
        ws_summary['A3'] = f'Generated: {timezone.now().strftime("%Y-%m-%d %H:%M:%S")}'
        
        # Statistics
        logs = Log.objects.filter(timestamp__range=[start_date, end_date])
        alerts = Alert.objects.filter(created_at__range=[start_date, end_date])
        
        ws_summary['A5'] = 'Metric'
        ws_summary['B5'] = 'Count'
        ws_summary['A5'].font = Font(bold=True)
        ws_summary['B5'].font = Font(bold=True)
        
        stats = [
            ('Total Logs', logs.count()),
            ('Critical Logs', logs.filter(level='CRITICAL').count()),
            ('Error Logs', logs.filter(level='ERROR').count()),
            ('Threats Detected', logs.filter(ai_threat_detected=True).count()),
            ('Total Alerts', alerts.count()),
            ('Critical Alerts', alerts.filter(severity='critical').count()),
            ('Resolved Alerts', alerts.filter(status='resolved').count()),
        ]
        
        for idx, (metric, count) in enumerate(stats, start=6):
            ws_summary[f'A{idx}'] = metric
            ws_summary[f'B{idx}'] = count
        
        # Logs Sheet
        ws_logs = wb.create_sheet('Logs')
        headers = ['Timestamp', 'Level', 'Service', 'Message', 'IP Address', 'Threat Detected']
        ws_logs.append(headers)
        
        for cell in ws_logs[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='1a1a2e', end_color='1a1a2e', fill_type='solid')
            cell.font = Font(bold=True, color='FFFFFF')
        
        for log in logs.order_by('-timestamp')[:1000]:
            ws_logs.append([
                log.timestamp.strftime('%Y-%m-%d %H:%M:%S'),
                log.level,
                log.service or '',
                log.message[:100],
                log.ip_address or '',
                'Yes' if log.ai_threat_detected else 'No'
            ])
        
        # Alerts Sheet
        ws_alerts = wb.create_sheet('Alerts')
        alert_headers = ['Created', 'Title', 'Severity', 'Status', 'IP Address']
        ws_alerts.append(alert_headers)
        
        for cell in ws_alerts[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color='e74c3c', end_color='e74c3c', fill_type='solid')
            cell.font = Font(bold=True, color='FFFFFF')
        
        for alert in alerts.order_by('-created_at')[:500]:
            ws_alerts.append([
                alert.created_at.strftime('%Y-%m-%d %H:%M:%S'),
                alert.title,
                alert.severity.upper(),
                alert.status.upper(),
                alert.ip_address or ''
            ])
        
        # Auto-size columns
        for ws in [ws_summary, ws_logs, ws_alerts]:
            for column in ws.columns:
                max_length = 0
                column = [cell for cell in column]
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(cell.value)
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column[0].column_letter].width = adjusted_width
        
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer
